from collections import Counter
import random

from openpyxl import Workbook, load_workbook
import pytest

from generate_negative_controls import (
    OUTPUT_HEADERS,
    generate_scramble,
    generate_workbook,
    main,
    matches_excluded_motif,
    reverse_sequence,
)


def test_scramble_preserves_composition_and_is_distinct():
    original = "ACDEFGHIKLMNPQRSTVWY"
    result = generate_scramble(original, random.Random(7))

    assert result.status == "OK"
    assert result.sequence != original
    assert Counter(result.sequence) == Counter(original)


def test_scramble_excludes_regex_match_anywhere():
    assert matches_excluded_motif("AALIPPAA")
    result = generate_scramble("AALIPPAA", random.Random(3))

    assert result.sequence is not None
    assert not matches_excluded_motif(result.sequence)


def test_reversal():
    assert reverse_sequence("MGS") == "SGM"


def test_scramble_is_reproducible_for_seed():
    sequence = "PVPATIPLSSVWVAENSD"
    first = generate_scramble(sequence, random.Random(1234))
    second = generate_scramble(sequence, random.Random(1234))

    assert first == second


def test_no_valid_scramble_for_repeated_single_residue():
    result = generate_scramble("AAAA", random.Random(1), max_attempts=5)

    assert result.sequence is None
    assert result.status.startswith("NO_VALID_SCRAMBLE")


def test_workbook_case_handling_names_flags_and_invalid_rows(tmp_path):
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "output.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("Name", "Sequence (AA, DNA, or RNA)"))
    sheet.append(("mixed", "mGs"))
    sheet.append(("motif_after_reverse", "AAPPILAA"))
    sheet.append(("missing", None))
    sheet.append(("invalid", "ACD*"))
    workbook.save(input_path)

    warnings = generate_workbook(input_path, output_path, seed=19)
    result = load_workbook(output_path, data_only=True)
    rows = list(result.active.iter_rows(values_only=True))

    assert rows[0] == OUTPUT_HEADERS
    assert rows[1][0:3] == ("mixed", "mGs", "MGS")
    assert rows[1][3] == "mixed_scrambled"
    assert rows[1][6:9] == ("mixed_reversed", "SGM", False)
    assert rows[2][7:9] == ("AALIPPAA", True)
    assert rows[3][4] is None
    assert rows[3][5].startswith("INVALID_INPUT")
    assert rows[4][4] is None
    assert any("row 4: missing sequence" in warning for warning in warnings)
    assert any("row 5: invalid amino-acid" in warning for warning in warnings)


def test_mixed_case_duplicates_abort_before_output_is_written(tmp_path, capsys):
    input_path = tmp_path / "duplicates.xlsx"
    output_path = tmp_path / "must-not-exist.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(("Name", "Sequence (AA, DNA, or RNA)"))
    sheet.append(("peptide-one", "mGs"))
    sheet.append(("unique", "ACDE"))
    sheet.append(("peptide-two", "MGS"))
    workbook.save(input_path)

    with pytest.raises(ValueError) as raised:
        generate_workbook(input_path, output_path, seed=19)

    message = str(raised.value)
    assert "duplicate input peptide sequences after normalization" in message
    assert "MGS" in message
    assert "row 2 (Name: 'peptide-one')" in message
    assert "row 4 (Name: 'peptide-two')" in message
    assert not output_path.exists()

    exit_status = main(
        [str(input_path), "--output", str(output_path), "--seed", "19"]
    )
    captured = capsys.readouterr()
    assert exit_status != 0
    assert "MGS" in captured.err
    assert "row 2 (Name: 'peptide-one')" in captured.err
    assert "row 4 (Name: 'peptide-two')" in captured.err
    assert not output_path.exists()
