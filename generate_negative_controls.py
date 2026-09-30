#!/usr/bin/env python3
"""Generate scrambled and reversed negative-control peptides from an Excel file."""

from __future__ import annotations

import argparse
import random
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, NamedTuple

from openpyxl import Workbook, load_workbook


NAME_HEADER = "Name"
SEQUENCE_HEADER = "Sequence (AA, DNA, or RNA)"
STANDARD_AMINO_ACIDS = frozenset("ACDEFGHIKLMNPQRSTVWY")
EXCLUDED_MOTIF = re.compile(
    r"(L[MLIV]PP)|(((L[LMIV]PA)|(L..P))[ILMVAFYW].[MLIVFHPAY])"
)

OUTPUT_HEADERS = (
    "Input Name",
    "Original Sequence",
    "Normalized Sequence",
    "Scrambled Control Name",
    "Scrambled Sequence",
    "Scrambled-Control Status",
    "Reversed Control Name",
    "Reversed Sequence",
    "Reversed Matches Regex",
)


class ScrambleResult(NamedTuple):
    sequence: str | None
    status: str


def normalize_sequence(value: Any) -> tuple[str | None, str | None]:
    """Return an uppercase standard-AA sequence and, on failure, an error."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, "missing sequence"
    if not isinstance(value, str):
        return None, f"sequence is not text ({type(value).__name__})"

    sequence = value.strip().upper()
    invalid = sorted(set(sequence) - STANDARD_AMINO_ACIDS)
    if invalid:
        rendered = ", ".join(repr(character) for character in invalid)
        return None, f"invalid amino-acid character(s): {rendered}"
    return sequence, None


def matches_excluded_motif(sequence: str) -> bool:
    """Search for the prohibited motif anywhere in a sequence."""
    return EXCLUDED_MOTIF.search(sequence) is not None


def reverse_sequence(sequence: str) -> str:
    return sequence[::-1]


def generate_scramble(
    sequence: str, rng: random.Random, max_attempts: int = 10_000
) -> ScrambleResult:
    """Find a bounded, composition-preserving, motif-free permutation.

    A valid candidate equal to the reversed sequence is retained only as a
    fallback while the remainder of the bounded search looks for a distinct
    alternative.
    """
    if max_attempts < 1:
        raise ValueError("max_attempts must be at least 1")
    if len(set(sequence)) < 2:
        return ScrambleResult(
            None,
            "NO_VALID_SCRAMBLE: sequence has no permutation different from itself",
        )

    reversed_sequence = reverse_sequence(sequence)
    reversed_fallback: str | None = None
    residues = list(sequence)

    for _ in range(max_attempts):
        rng.shuffle(residues)
        candidate = "".join(residues)
        if candidate == sequence or matches_excluded_motif(candidate):
            continue
        if candidate == reversed_sequence:
            reversed_fallback = candidate
            continue
        return ScrambleResult(candidate, "OK")

    if reversed_fallback is not None:
        return ScrambleResult(
            reversed_fallback,
            "OK_REVERSED_FALLBACK: no other valid scramble found within search limit",
        )
    return ScrambleResult(
        None,
        f"NO_VALID_SCRAMBLE: none found in {max_attempts} bounded attempts",
    )


def _header_map(worksheet: Any) -> dict[str, int]:
    headers = {
        value: index
        for index, value in enumerate(
            next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        )
        if isinstance(value, str)
    }
    missing = [
        header for header in (NAME_HEADER, SEQUENCE_HEADER) if header not in headers
    ]
    if missing:
        raise ValueError(f"missing required column(s): {', '.join(missing)}")
    return headers


def _reject_duplicate_sequences(
    rows: list[tuple[Any, ...]], headers: dict[str, int]
) -> None:
    """Reject valid sequences repeated after normalizing case and whitespace."""
    occurrences: dict[str, list[tuple[int, Any]]] = defaultdict(list)
    name_column = headers[NAME_HEADER]
    sequence_column = headers[SEQUENCE_HEADER]

    for excel_row, row in enumerate(rows, start=2):
        name = row[name_column] if name_column < len(row) else None
        original = row[sequence_column] if sequence_column < len(row) else None
        normalized, error = normalize_sequence(original)
        if error is None:
            assert normalized is not None
            occurrences[normalized].append((excel_row, name))

    duplicates = {
        sequence: locations
        for sequence, locations in occurrences.items()
        if len(locations) > 1
    }
    if not duplicates:
        return

    details = []
    for sequence, locations in duplicates.items():
        rows_and_names = ", ".join(
            f"row {row_number} (Name: {name!r})"
            for row_number, name in locations
        )
        details.append(f"  {sequence}: {rows_and_names}")
    raise ValueError(
        "duplicate input peptide sequences after normalization:\n"
        + "\n".join(details)
    )


def generate_workbook(
    input_path: Path, output_path: Path, seed: int, max_attempts: int = 10_000
) -> list[str]:
    """Read the active sheet, write controls, and return row-level warnings."""
    source = load_workbook(input_path, read_only=True, data_only=True)
    try:
        worksheet = source.active
        headers = _header_map(worksheet)
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    finally:
        source.close()

    # Complete duplicate validation before constructing controls or touching
    # the requested output path.
    _reject_duplicate_sequences(rows, headers)

    output = Workbook()
    output_sheet = output.active
    output_sheet.title = "Negative controls"
    output_sheet.append(OUTPUT_HEADERS)
    warnings: list[str] = []
    rng = random.Random(seed)

    for excel_row, row in enumerate(rows, start=2):
        name_value = row[headers[NAME_HEADER]] if headers[NAME_HEADER] < len(row) else None
        original = (
            row[headers[SEQUENCE_HEADER]]
            if headers[SEQUENCE_HEADER] < len(row)
            else None
        )
        name = "" if name_value is None else str(name_value)
        scrambled_name = f"{name}_scrambled" if name else ""
        reversed_name = f"{name}_reversed" if name else ""

        if not name:
            warnings.append(f"row {excel_row}: missing Name")

        normalized, error = normalize_sequence(original)
        if error is not None:
            warnings.append(f"row {excel_row}: {error}")
            output_sheet.append(
                (
                    name_value,
                    original,
                    None,
                    scrambled_name,
                    None,
                    f"INVALID_INPUT: {error}",
                    reversed_name,
                    None,
                    False,
                )
            )
            continue

        assert normalized is not None
        scramble = generate_scramble(normalized, rng, max_attempts)
        reversed_sequence = reverse_sequence(normalized)
        if scramble.sequence is None:
            warnings.append(f"row {excel_row}: {scramble.status}")
        output_sheet.append(
            (
                name_value,
                original,
                normalized,
                scrambled_name,
                scramble.sequence,
                scramble.status,
                reversed_name,
                reversed_sequence,
                matches_excluded_motif(reversed_sequence),
            )
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output.save(output_path)
    return warnings


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate scrambled and reversed peptide controls in XLSX format."
    )
    parser.add_argument("input", type=Path, help="input .xlsx workbook")
    parser.add_argument("--output", type=Path, required=True, help="output .xlsx path")
    parser.add_argument("--seed", type=int, required=True, help="random seed")
    parser.add_argument(
        "--max-attempts",
        type=int,
        default=10_000,
        help="maximum shuffle attempts per peptide (default: 10000)",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.max_attempts < 1:
        print("error: --max-attempts must be at least 1", file=sys.stderr)
        return 2
    try:
        warnings = generate_workbook(
            args.input, args.output, args.seed, args.max_attempts
        )
    except (OSError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 1

    print(f"Wrote {args.output}")
    if warnings:
        print("Rows requiring attention:", file=sys.stderr)
        for warning in warnings:
            print(f"  {warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
